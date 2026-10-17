from __future__ import annotations

import io
import re

from pydantic import BaseModel, ConfigDict, Field

from agent.contracts import SourceSpan


class ParsedSegment(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: str
    text: str
    span: SourceSpan
    metadata: dict[str, str] = Field(default_factory=dict)


class ParsedDocument(BaseModel):
    model_config = ConfigDict(extra="forbid")
    source_doc_id: str
    text: str
    segments: list[ParsedSegment]
    parser_tier: str
    degraded: bool = False
    degraded_reason: str | None = None


def _records(text: str, source_doc_id: str, confidence: float = 0.9) -> list[ParsedSegment]:
    segments: list[ParsedSegment] = []
    for match in re.finditer(r"[^\n]+", text):
        line = match.group(0).strip()
        if not line:
            continue
        lower = line.casefold()
        kind = "experience" if any(word in lower for word in ("intern", "engineer", "developer", "analyst")) else "text"
        if re.search(r"\b(project|built|developed|implemented|designed)\b", lower):
            kind = "project"
        if re.search(r"\b(course|coursework|semester|credits|grade|cgpa|gpa)\b", lower):
            kind = "course"
        segments.append(ParsedSegment(
            kind=kind,
            text=line,
            span=SourceSpan(
                source_doc_id=source_doc_id,
                char_start=match.start(),
                char_end=match.end(),
                extractor_confidence=confidence,
                quote=line[:500],
            ),
        ))
    return segments


def _pdf_text(file_bytes: bytes) -> tuple[str, str, bool, str | None]:
    try:
        import fitz

        document = fitz.open(stream=file_bytes, filetype="pdf")
        page_text: list[str] = []
        used_ocr = False
        for page in document:
            blocks = [block for block in page.get_text("blocks") if len(block) > 4 and str(block[4]).strip()]
            if blocks:
                # Group blocks by horizontal column before sorting top-to-bottom.
                midpoint = page.rect.width / 2
                left = sorted((block for block in blocks if block[0] < midpoint), key=lambda row: (row[1], row[0]))
                right = sorted((block for block in blocks if block[0] >= midpoint), key=lambda row: (row[1], row[0]))
                ordered = left + right if left and right else sorted(blocks, key=lambda row: (row[1], row[0]))
                page_text.append("\n".join(str(block[4]).strip() for block in ordered))
                continue
            try:
                import pytesseract
                from PIL import Image

                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                page_text.append(pytesseract.image_to_string(image))
                used_ocr = True
            except (ImportError, RuntimeError, OSError):
                page_text.append("")
        text = "\n".join(page_text).strip()
        if not text:
            return "", "pdf", True, "PDF had no extractable text and OCR was unavailable or produced no text."
        return text, "pdf_ocr" if used_ocr else "pdf_layout", used_ocr, None
    except ImportError:
        return "", "pdf", True, "PyMuPDF is not installed; PDF parsing is unavailable."


def _docx_text(file_bytes: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _kind_from_filename(filename: str) -> str:
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else "text"
    return suffix


def parse_document(file_bytes: bytes, kind: str, source_doc_id: str = "profile") -> ParsedDocument:
    """Extract layout-aware text and line segments with source character offsets."""
    file_kind = kind.casefold().lstrip(".")
    degraded = False
    reason: str | None = None
    if file_kind == "pdf":
        text, tier, degraded, reason = _pdf_text(file_bytes)
    elif file_kind == "docx":
        try:
            text = _docx_text(file_bytes)
            tier = "docx_tables"
        except ImportError:
            text, tier, degraded, reason = "", "docx", True, "python-docx is not installed."
    else:
        text = file_bytes.decode("utf-8", errors="replace")
        tier = "text"
    segments = _records(text, source_doc_id, confidence=0.78 if tier == "pdf_ocr" else 0.94)
    return ParsedDocument(
        source_doc_id=source_doc_id,
        text=text,
        segments=segments,
        parser_tier=tier,
        degraded=degraded,
        degraded_reason=reason,
    )


def parse_text(text: str, source_doc_id: str = "profile") -> ParsedDocument:
    encoded = text.encode("utf-8")
    return parse_document(encoded, "txt", source_doc_id)
